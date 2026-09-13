"""One-time local import. Never writes into the public archive or overwrites tracker edits."""
import json, sys, hashlib
from pathlib import Path
from datetime import datetime
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'.local-work-tracker'/'records.json'

def text(value):
    return value.isoformat()[:10] if isinstance(value,datetime) else str(value).strip() if value is not None else ''

def base(category,title,sheet,row,cells,stage='Needs confirmation'):
    return dict(id='WORK-'+hashlib.sha256(f'{sheet}:{row}:{title}'.encode()).hexdigest()[:10].upper(),category=category,title=title,stage=stage,owner='',dueDate='',nextAction='',notes='',workingUrl='',publishedUrl='',recipient='',sentDate='',response='',printStatus='Not planned',websiteStatus='Not planned',distributionStatus='Not planned',flags=['Imported status needs team confirmation'],source=dict(sheet=sheet,row=row,cells=cells),revision=1,updatedAt='',history=[dict(at='',summary=f'Imported from Media.xlsx · {sheet}, row {row}')])

def main():
    if OUT.exists():raise SystemExit('Import stopped: local records already exist. Existing edits were preserved.')
    book=openpyxl.load_workbook(sys.argv[1],data_only=False);items=[]
    mapping={'Design':'Design','Design Stage':'Design','Content stage':'Content preparation','Content  generation':'Content preparation','Content Generation':'Content preparation','Review Stage':'Review','Concept stage':'Concept','Editing Stage':'Editing','Proposed':'Concept'}
    for sheet in book:
        if sheet.title=='Annual Report':
            cells={c.coordinate:text(c.value) for row in sheet for c in row if c.value is not None}
            item=base('Annual reports','Annual report — year to confirm',sheet.title,1,cells,'Proofreading');item['notes']='Draft designed and ready: '+text(sheet['C1'].value)+'\nCurrent status: '+text(sheet['C2'].value);item['flags'].append('Confirm report year and current proofreading owner');items.append(item);continue
        if sheet.title=='Sheet1':
            for col in range(2,sheet.max_column+1):
                title=text(sheet.cell(1,col).value)
                if not title:continue
                recipients=[sheet.cell(row,col) for row in range(2,sheet.max_row+1) if sheet.cell(row,col).value]
                for cell in recipients or [sheet.cell(2,col)]:
                    item=base('Distribution',title,sheet.title,cell.row,{sheet.cell(1,col).coordinate:title,cell.coordinate:text(cell.value)},'Planned');item['recipient']=text(cell.value)
                    if not item['recipient']:item['flags'].append('Recipient not recorded')
                    items.append(item)
            continue
        for row in range(2,sheet.max_row+1):
            cells={c.coordinate:text(c.value) for c in sheet[row] if c.value is not None}
            if not cells:continue
            values=[text(sheet.cell(row,c).value) for c in range(1,8)]
            if sheet.title=='Publications':
                if not values[1]:continue
                stage='On hold' if values[3].lower()=='on hold' else mapping.get(values[2],'Needs confirmation')
                item=base('Publications',values[1],sheet.title,row,cells,stage)
                item['notes']=f'Original status: {values[2]}\nDissemination / progress: {values[3]}'
                if values[2]=='Yes':item['flags'].append('“Yes” does not establish printing, upload or distribution completion')
                if values[4]:item['flags'].append('Unlabelled source date: '+values[4]+' — confirm its meaning before setting a deadline')
            elif sheet.title=='News':
                item=base('News & press releases',values[2] or values[1],sheet.title,row,cells,'Proposed' if row in range(6,11) else 'Needs confirmation')
                item['notes']=f'Type: {values[1]}\nPublisher: {values[3]}\nPage: {values[4]}\nLanguage: {values[5]}\nDG quote / source note: {values[6]}'
                item['flags'].append('Publication date and source evidence need confirmation')
                if values[0]>'2026-09-11' and isinstance(sheet.cell(row,1).value,datetime):item['flags'].append('Stored publication date is in the future; possible day/month reversal')
            elif sheet.title=='Sampada':
                status={'Done':'Needs confirmation','Design':'Design','Content generation':'Content preparation','Planned':'Planning','Planed':'Planning','to be discussed':'Planning'}.get(values[2],'Planning')
                item=base('Sampada',f'Sampada — {values[0]}'+(f' · {values[1]}' if values[1] else ''),sheet.title,row,cells,status)
                item['notes']=f'Original status: {values[2]}\nPosting: {values[3]}'+(f'\nUnlabelled theme note: {values[5]}' if values[5] else '')
                item['flags'].append('Issue year is not recorded')
            elif sheet.title=='Representations':
                item=base('Representations',values[1],sheet.title,row,cells,'Sent' if values[5]=='Sent' else 'Needs confirmation')
                item['recipient']=values[3]+(' · '+values[4] if values[4] else '')
                item['notes']=f'Department: {values[2]}\nOriginal date: {values[0]}';item['response']=values[6]
                item['nextAction']='Confirm acknowledgement and next follow-up'
            else:continue
            items.append(item)
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(items,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({category:sum(x['category']==category for x in items) for category in dict.fromkeys(x['category'] for x in items)}))
if __name__=='__main__':main()
